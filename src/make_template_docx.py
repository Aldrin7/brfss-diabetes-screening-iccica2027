"""Build the manuscript on the conference's own Word template (conference-template-a4.docx).

The supplied template is a Strict-OOXML Word file, so it is first rewritten into the
transitional namespace that python-docx can open; its styles.xml, section setup, footer
and numbering are then reused verbatim. The manuscript content comes from manuscript.tex
and is emitted with the template's own style names:

    paper title, Author, Abstract, Keywords,
    Heading 1/2, Body Text,
    table head, table col head, table copy, table footnote, figure caption,
    Heading 5, references

Usage:
    python src/make_template_docx.py                 # -> manuscript_template.docx
    python src/make_template_docx.py --out x.docx
"""
import argparse
import os
import re
import shutil
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_docx import (  # noqa: E402
    extract_bibliography, extract_command_arg, inline_math_to_text, math_to_text,
    normalize_multicolumn, number_citations,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(ROOT, "build")
TEMPLATE = os.path.join(ROOT, "conference-template-a4.docx")
TEX = os.path.join(ROOT, "manuscript.tex")
FIGS = [os.path.join(ROOT, "figures"), os.path.join(ROOT, "results", "figures")]

NS_MAP = [
    ("http://purl.oclc.org/ooxml/officeDocument/relationships",
     "http://schemas.openxmlformats.org/officeDocument/2006/relationships"),
    ("http://purl.oclc.org/ooxml/officeDocument/extendedProperties",
     "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"),
    ("http://purl.oclc.org/ooxml/officeDocument/customXml",
     "http://schemas.openxmlformats.org/officeDocument/2006/customXml"),
    ("http://purl.oclc.org/ooxml/officeDocument/docPropsVTypes",
     "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"),
    ("http://purl.oclc.org/ooxml/officeDocument/math",
     "http://schemas.openxmlformats.org/officeDocument/2006/math"),
    ("http://purl.oclc.org/ooxml/wordprocessingml/main",
     "http://schemas.openxmlformats.org/wordprocessingml/2006/main"),
    ("http://purl.oclc.org/ooxml/drawingml/main",
     "http://schemas.openxmlformats.org/drawingml/2006/main"),
    ("http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing",
     "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"),
]

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
ALPHA = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


# --------------------------------------------------------------------------- template
def strict_to_transitional(src, dst):
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                txt = data.decode("utf-8", "replace")
                for a, b in NS_MAP:
                    txt = txt.replace(a, b)
                data = txt.encode("utf-8")
            zout.writestr(item, data)
    zin.close()
    return dst


# --------------------------------------------------------------------------- parsing
def parse_tex(text):
    """Return a list of blocks: (kind, payload)."""
    blocks = []
    lines = text.split("\n")
    # offset of the start of each line, used to skip multi-line commands
    starts, off = [], 0
    for l in lines:
        starts.append(off)
        off += len(l) + 1

    def line_of(pos):
        lo, hi = 0, len(starts) - 1
        while lo < hi:
            mid = (lo + hi + 1) // 2
            if starts[mid] <= pos:
                lo = mid
            else:
                hi = mid - 1
        return lo

    i = 0
    in_body = False
    while i < len(lines):
        line = lines[i].strip()
        if not in_body:
            if line.startswith("\\begin{document}"):
                in_body = True
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("\\title{") or line.startswith("\\author{"):
            cmd = "title" if line.startswith("\\title{") else "author"
            arg, span = extract_command_arg(text, cmd)
            blocks.append((("authors" if cmd == "author" else "title"), arg))
            i = line_of(span[1]) + 1
            continue
        if line.startswith("\\begin{abstract}"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\end{abstract}"):
                body.append(lines[i].strip())
                i += 1
            i += 1
            blocks.append(("abstract", " ".join(body)))
            continue
        if line.startswith("\\begin{IEEEkeywords}"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\end{IEEEkeywords}"):
                body.append(lines[i].strip())
                i += 1
            i += 1
            blocks.append(("keywords", " ".join(body)))
            continue
        if line.startswith("\\section{"):
            m = re.match(r"\\section\{(.*)\}", line)
            blocks.append(("h1", m.group(1)))
            i += 1
            continue
        if line.startswith("\\section*{"):
            m = re.match(r"\\section\*\{(.*)\}", line)
            blocks.append(("h1*", m.group(1)))
            i += 1
            continue
        if line.startswith("\\subsection{"):
            m = re.match(r"\\subsection\{(.*)\}", line)
            blocks.append(("h2", m.group(1)))
            i += 1
            continue
        if line.startswith("\\begin{table"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\end{table"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("table", "\n".join(body)))
            continue
        if line.startswith("\\begin{figure"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\end{figure"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("figure", "\n".join(body)))
            continue
        if line.startswith("\\bibitem{"):
            m = re.match(r"\\bibitem\{(\w+)\}\s*(.*)", line)
            key, first = m.group(1), m.group(2)
            body = [first] if first.strip() else []
            i += 1
            # the entry text may continue on the following lines
            while i < len(lines):
                nxt = lines[i].strip()
                if (not nxt or nxt.startswith("\\bibitem")
                        or nxt.startswith("\\end{thebibliography}")
                        or nxt.startswith("\\begin{") or nxt.startswith("\\end{")):
                    break
                body.append(nxt)
                i += 1
            blocks.append(("ref", (key, " ".join(body))))
            continue
        if line.startswith("\\begin{thebibliography}"):
            i += 1
            continue
        if line.startswith("\\end{thebibliography}"):
            i += 1
            continue
        if line.startswith("\\end{document}"):
            break
        if line:
            blocks.append(("body", line))
        i += 1
    return blocks


def clean_text(s, strip=True):
    """LaTeX -> plain text for a paragraph (math already converted)."""
    s = s.replace("\\%", "%").replace("\\&", "&").replace("\\_", "_")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = s.replace("~", " ").replace("\\$", "$")
    s = re.sub(r"\\textit\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\texttt\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\emph\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\[a-zA-Z]+\s*", "", s)
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip() if strip else s


def segments(text):
    """Split a paragraph into (text, bold) runs, honouring \\textbf{}.

    Interior spacing is preserved: a bold lead-in keeps the space that follows it."""
    parts = []
    rest = text
    while True:
        m = re.search(r"\\textbf\{([^{}]*)\}", rest)
        if not m:
            parts.append((rest, False))
            break
        if m.start():
            parts.append((rest[:m.start()], False))
        parts.append((m.group(1), True))
        rest = rest[m.end():]
    out = []
    for i, (t, b) in enumerate(parts):
        t = clean_text(t, strip=False)
        if i == 0:
            t = t.lstrip()
        if i == len(parts) - 1:
            t = t.rstrip()
        if t:
            out.append((t, b))
    return out


def balanced(text, start):
    """Return the balanced {...} content beginning at text[start] == '{'."""
    depth, i = 0, start
    while i < len(text):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:i], i + 1
        i += 1
    return text[start + 1:], len(text)


def parse_table(body, labels):
    """Parse a LaTeX table into raw (uncleaned) caption and cells.

    The text is left raw so that the builder can run the math converter over it
    before the final LaTeX cleanup."""
    cap = re.search(r"\\caption\{(.*?)\}\s*\\label", body, re.S)
    lab = re.search(r"\\label\{([^}]*)\}", body)
    caption = cap.group(1) if cap else ""
    tab = re.search(r"\\begin\{tabular\}(.*?)\\end\{tabular\}", body, re.S)
    rows = []
    if tab:
        raw = normalize_multicolumn(tab.group(0)).replace("{@{}l}", "l")
        raw = re.sub(r"\\(toprule|midrule|bottomrule|hline)", "", raw)
        raw = re.sub(r"\\begin\{tabular\}\{[^}]*\}", "", raw)
        raw = raw.replace("\\end{tabular}", "")
        for rline in raw.split("\\\\"):
            rline = rline.strip()
            if not rline:
                continue
            cells = [c.strip() for c in rline.split("&")]
            if not any(cells):
                continue
            rows.append(cells)
    return {"caption": caption, "rows": rows, "label": lab.group(1) if lab else None}


# --------------------------------------------------------------------------- builder
def add_runs(p, text, bold_default=False):
    for seg, bold in segments(text):
        r = p.add_run(seg)
        r.bold = bold or bold_default
    return p


def set_borders(table, doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def edge(tag, val, sz=6):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        return el

    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom"):
        borders.append(edge("w:%s" % side, "single", 6))
    for side in ("left", "right", "insideH", "insideV"):
        borders.append(edge("w:%s" % side, "none", 0))
    tblPr.append(borders)
    # rule under the header row
    if table.rows:
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcB = OxmlElement("w:tcBorders")
            tcB.append(edge("w:bottom", "single", 6))
            tcPr.append(tcB)


def set_columns(table, widths_in):
    """Fixed column widths (inches) and repeat-header / no-row-split settings."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Inches

    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_in):
                cell.width = Inches(widths_in[i])
    for i, col in enumerate(table.columns):
        if i < len(widths_in):
            col.width = Inches(widths_in[i])
    # repeat the header row on page breaks, never split a row
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))


COLUMN_WIDTHS = {
    4: [1.45, 1.35, 1.35, 1.35],
    5: [1.45, 1.0, 1.0, 1.0, 1.0],
    6: [1.45, 0.8, 0.8, 0.8, 0.8, 0.8],
}


def build(out_path):
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    for d in FIGS:
        if os.path.isdir(d):
            for f in os.listdir(d):
                if f.lower().endswith(".png"):
                    shutil.copy2(os.path.join(d, f), os.path.join(BUILD, f))

    tpl = strict_to_transitional(TEMPLATE, os.path.join(BUILD, "template_transitional.docx"))
    doc = docx.Document(tpl)

    # keep the template's section setup, drop its guidance content
    body = doc.element.body
    sectPr = body.find(docx.oxml.ns.qn("w:sectPr"))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)

    text = open(TEX, encoding="utf-8").read()
    items, _ = extract_bibliography(text)
    blocks = parse_tex(text)

    # resolve citations and cross references
    order = {k: i + 1 for i, (k, _) in enumerate(items)}
    label_num = {}
    tnum = fnum = 0
    for kind, payload in blocks:
        if kind == "table":
            tnum += 1
            tb = parse_table(payload, label_num)
            if tb["label"]:
                label_num[tb["label"]] = tnum
        elif kind == "figure":
            fnum += 1
            lab = re.search(r"\\label\{([^}]*)\}", payload)
            if lab:
                label_num[lab.group(1)] = fnum

    def cite(m):
        nums = sorted(order[k.strip()] for k in m.group(1).split(",") if k.strip() in order)
        if not nums:
            return "[]"
        if len(nums) == 1:
            return "[%d]" % nums[0]
        if nums == list(range(nums[0], nums[0] + len(nums))):
            return "[%d]-[%d]" % (nums[0], nums[-1])
        return "[" + ",".join(str(n) for n in nums) + "]"

    def ref(m):
        return str(label_num.get(m.group(1), "?"))

    def prep(s):
        s = number_citations(s, items)
        s = re.sub(r"\\ref\{([^}]*)\}", ref, s)
        s = inline_math_to_text(s)
        s = s.replace("\\%", "%")
        return s

    sec = 0
    sub = 0
    tbl_i = 0
    fig_i = 0
    for kind, payload in blocks:
        if kind == "title":
            p = doc.add_paragraph(style="paper title")
            add_runs(p, clean_text(payload))
        elif kind == "authors":
            a = payload
            m = re.search(r"\\IEEEauthorblockN\{", a)
            if m:
                content, end = balanced(a, m.end() - 1)
                a = a[:m.start()] + content + a[end:]
            m = re.search(r"\\IEEEauthorblockA\{", a)
            if m:
                content, end = balanced(a, m.end() - 1)
                a = a[:m.start()] + content + a[end:]
            a = a.replace("\\\\", "\n")
            a = re.sub(r"\\textit\{([^{}]*)\}", r"\1", a)
            a = a.replace("{", "").replace("}", "").replace("\\", "")
            for line in [x.strip() for x in a.split("\n") if x.strip()]:
                p = doc.add_paragraph(style="Author")
                add_runs(p, line)
        elif kind == "abstract":
            p = doc.add_paragraph(style="Abstract")
            add_runs(p, "Abstract\u2014" + prep(payload))
        elif kind == "keywords":
            p = doc.add_paragraph(style="Keywords")
            add_runs(p, "Keywords\u2014" + prep(payload))
        elif kind == "h1":
            sec += 1
            sub = 0
            # the template's Heading 1 style carries the automatic Roman numbering
            p = doc.add_paragraph(style="Heading 1")
            add_runs(p, clean_text(prep(payload)))
        elif kind == "h1*":
            p = doc.add_paragraph(style="Heading 5")
            add_runs(p, clean_text(prep(payload)))
        elif kind == "h2":
            sub += 1
            # the template's Heading 2 style carries the automatic lettered numbering
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, clean_text(prep(payload)))
        elif kind == "body":
            txt = clean_text(prep(payload))
            if not txt:
                continue
            p = doc.add_paragraph(style="Body Text")
            add_runs(p, prep(payload))
        elif kind == "table":
            tbl_i += 1
            tb = parse_table(payload, label_num)
            if tb["caption"]:
                p = doc.add_paragraph(style="table head")
                add_runs(p, prep(tb["caption"]))
            rows = tb["rows"]
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell = t.cell(ri, ci)
                        val = row[ci] if ci < len(row) else ""
                        cp = cell.paragraphs[0]
                        cp.style = doc.styles["table col head" if ri == 0 else "table copy"]
                        add_runs(cp, prep(val))
                set_borders(t, doc)
                set_columns(t, COLUMN_WIDTHS.get(ncols, [5.4 / ncols] * ncols))
                # group rows (label in the first cell, the rest empty) span the full width
                for ri, row in enumerate(rows):
                    if ri == 0:
                        continue
                    if row and row[0].strip() and not any(c.strip() for c in row[1:]):
                        t.cell(ri, 0).merge(t.cell(ri, ncols - 1))
        elif kind == "figure":
            fig_i += 1
            img = re.search(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}", payload)
            cap = re.search(r"\\caption\{(.*?)\}\s*\\label", payload, re.S)
            if img:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                path = os.path.join(BUILD, img.group(1))
                if not os.path.exists(path):
                    path = os.path.join(ROOT, "figures", img.group(1))
                p.add_run().add_picture(path, width=Inches(4.2))
            if cap:
                p = doc.add_paragraph(style="figure caption")
                add_runs(p, prep(cap.group(1)))
        elif kind == "ref":
            key, bodytxt = payload
            p = doc.add_paragraph(style="references")
            add_runs(p, prep(bodytxt))

    doc.save(out_path)
    return out_path, blocks


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="manuscript_template.docx")
    args = ap.parse_args()
    out = os.path.join(ROOT, args.out)
    path, blocks = build(out)
    import docx
    d = docx.Document(path)
    print("wrote %s" % path)
    print("  paragraphs: %d | tables: %d | images: %d"
          % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))
    styles = {}
    for p in d.paragraphs:
        styles[p.style.name] = styles.get(p.style.name, 0) + 1
    for k in sorted(styles):
        print("   %-16s %d" % (k, styles[k]))


if __name__ == "__main__":
    main()
