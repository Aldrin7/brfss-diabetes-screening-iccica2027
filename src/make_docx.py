"""Convert manuscript.tex to a Word (.docx) file for co-author review and for venues
that accept a Word submission.

The PDF from src/compile_ieee.py remains the authoritative camera-ready artifact; this
converter reproduces the same content, tables, figure, numbered citations and references
in Word, with every number rendered as selectable text (inline LaTeX math is converted to
plain Unicode) so the document can be read, searched and annotated.

Pipeline:
  1. read manuscript.tex, keep only the body
  2. resolve \\cite{...} to IEEE-style numbers using the \\bibitem order
  3. rewrite the thebibliography environment as a numbered reference list
  4. convert tables/figures into caption paragraph + tabular/includegraphics, numbering the
     captions as they appear in the PDF
  5. convert inline math to plain text with Unicode symbols
  6. run pandoc (latex -> docx), embedding the figure
  7. post-process with python-docx: Times New Roman 10pt, 0.75in margins, 1 or 2 columns,
     first-line indent, caption styling

Usage:
    python src/make_docx.py            # -> manuscript.docx (single column, editable)
    python src/make_docx.py --twocol   # -> manuscript_2col.docx (matches the PDF layout)
"""
import argparse
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(ROOT, "build")
TEX = os.path.join(ROOT, "manuscript.tex")
FIGS = [os.path.join(ROOT, "figures"), os.path.join(ROOT, "results", "figures")]

SYMBOLS = [
    (r"\{,\}", ","), (r"\%", "%"), (r"\times", "\u00d7"), (r"\pm", "\u00b1"),
    (r"\ge", "\u2265"), (r"\le", "\u2264"), (r"\rightarrow", "\u2192"),
    (r"\Delta", "\u0394"), (r"\sigma", "\u03c3"), (r"\alpha", "\u03b1"),
    (r"\beta", "\u03b2"), (r"\theta", "\u03b8"), (r"\lambda", "\u03bb"),
    (r"\mu", "\u03bc"), (r"\pi", "\u03c0"), (r"\tau", "\u03c4"),
    (r"\phi", "\u03c6"), (r"\epsilon", "\u03b5"), (r"\eta", "\u03b7"),
    (r"\rho", "\u03c1"), (r"\ldots", "\u2026"), (r"\ ", " "), (r"\,", ""),
    (r"\;", " "), (r"\:", ""), (r"\!", ""), (r"~", " "),
]


def stage_figures():
    os.makedirs(BUILD, exist_ok=True)
    for d in FIGS:
        if os.path.isdir(d):
            for f in os.listdir(d):
                if f.lower().endswith(".png"):
                    shutil.copy2(os.path.join(d, f), os.path.join(BUILD, f))


def math_to_text(s):
    for a, b in SYMBOLS:
        s = s.replace(a, b)
    s = re.sub(r"\\(?:text|emph|mathrm|texttt|mathbf|operatorname)\{([^{}]*)\}", r"\1", s)
    s = s.replace("\\left", "").replace("\\right", "")
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def inline_math_to_text(text):
    def repl(m):
        return math_to_text(m.group(1))
    # display math is kept as an equation, inline math becomes text
    text = re.sub(r"\$\$(.+?)\$\$", lambda m: "$$%s$$" % m.group(1), text, flags=re.S)
    text = re.sub(r"\\\[(.+?)\\\]", lambda m: "$$%s$$" % m.group(1), text, flags=re.S)
    text = re.sub(r"(?<!\$)\$([^$\n]+?)\$(?!\$)", repl, text)
    return text


def extract_bibliography(text):
    m = re.search(r"\\begin\{thebibliography\}\{[^}]*\}(.*?)\\end\{thebibliography\}", text, re.S)
    if not m:
        return [], ""
    items = []
    for bm in re.finditer(r"\\bibitem\{(\w+)\}\s*(.*?)(?=\\bibitem|\Z)", m.group(1), re.S):
        key, body = bm.group(1), bm.group(2)
        body = re.sub(r"\\url\{([^}]*)\}", r"\1", body)
        body = re.sub(r"\\textit\{([^}]*)\}", r"\1", body)
        body = re.sub(r"\\texttt\{([^}]*)\}", r"\1", body)
        body = re.sub(r"``([^`]*)''", r"\1", body)
        body = re.sub(r"\\[a-zA-Z]+\s*", "", body)
        body = body.replace("{", "").replace("}", "").replace("\\", "")
        body = re.sub(r"\s+", " ", body).strip()
        items.append((key, body))
    return items, m.group(0)


def number_citations(text, items):
    order = {k: i + 1 for i, (k, _) in enumerate(items)}

    def repl(m):
        nums = sorted(order[k.strip()] for k in m.group(1).split(",") if k.strip() in order)
        if not nums:
            return "[]"
        if len(nums) == 1:
            return "[%d]" % nums[0]
        if nums == list(range(nums[0], nums[0] + len(nums))):
            return "[%d]-[%d]" % (nums[0], nums[-1])
        return "[" + ",".join(str(n) for n in nums) + "]"

    return re.sub(r"\\cite\{([^{}]*)\}", repl, text)


def normalize_multicolumn(tab):
    """Turn \\multicolumn{n}{spec}{text} rows into plain rows so pandoc can read them."""
    def repl(m):
        n, body = int(m.group(1)), m.group(2)
        body = re.sub(r"\\textit\{([^}]*)\}", r"\1", body)
        body = body.replace("\\", "").replace("{", "").replace("}", "").strip()
        cells = [body] + [""] * (n - 1)
        return " & ".join(cells) + " \\\\"
    return re.sub(r"\\multicolumn\{(\d+)\}\{(?:[^{}]|\{[^{}]*\})*\}\{((?:[^{}]|\{[^{}]*\})*)\}(?=\s*\\\\)", repl, tab)


def convert_tables_and_figures(text):
    """Return text with each table/figure replaced by a caption paragraph + content,
    plus a map of LaTeX labels to their printed numbers."""
    counters = {"table": 0, "figure": 0}
    labels = {}

    def table_repl(m):
        counters["table"] += 1
        n = counters["table"]
        body = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", body)
        if lab:
            labels[lab.group(1)] = ("Table", n)
        cap = re.search(r"\\caption\{(.*?)\}\s*\\label", body, re.S)
        caption = cap.group(1) if cap else ""
        caption = re.sub(r"\\emph\{([^}]*)\}", r"\1", caption)
        caption = re.sub(r"\\[a-zA-Z]+\s*", "", caption).replace("{", "").replace("}", "")
        caption = caption.strip()
        tab = re.search(r"\\begin\{tabular\}.*?\\end\{tabular\}", body, re.S)
        out = ""
        if caption:
            out += "\n\nTABLECAP%d %s\n\n" % (n, caption)
        if tab:
            t = normalize_multicolumn(tab.group(0))
            t = t.replace("{@{}l}", "l")
            t = re.sub(r"\\(toprule|midrule|bottomrule|hline)", "", t)
            out += t + "\n\n"
        return out

    def figure_repl(m):
        counters["figure"] += 1
        n = counters["figure"]
        body = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", body)
        if lab:
            labels[lab.group(1)] = ("Fig.", n)
        cap = re.search(r"\\caption\{(.*?)\}\s*\\label", body, re.S)
        caption = cap.group(1) if cap else ""
        caption = re.sub(r"\\emph\{([^}]*)\}", r"\1", caption)
        caption = re.sub(r"\\[a-zA-Z]+\s*", "", caption).replace("{", "").replace("}", "")
        caption = caption.strip()
        img = re.search(r"\\includegraphics\[([^\]]*)\]\{([^}]*)\}", body)
        out = ""
        if img:
            width = img.group(1).replace(" ", "")
            m2 = re.search(r"width=([\d.]+)\\linewidth", width)
            if m2:
                width = "width=%s\\linewidth" % m2.group(1)
            out += "\n\n\\includegraphics[%s]{%s}\n\n" % (width, img.group(2))
        if caption:
            out += "FIGCAP%d %s\n\n" % (n, caption)
        return out

    text = re.sub(r"\\begin\{table\*?\}(.*?)\\end\{table\*?\}", table_repl, text, flags=re.S)
    text = re.sub(r"\\begin\{figure\*?\}(.*?)\\end\{figure\*?\}", figure_repl, text, flags=re.S)
    return text, counters, labels


def extract_command_arg(text, cmd):
    """Return the brace-balanced argument of \\cmd{...} and its span."""
    m = re.search(r"\\" + cmd + r"\{", text)
    if not m:
        return None, None
    i = m.end()
    depth, start = 1, i
    while i < len(text) and depth:
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
        i += 1
    return text[start:i - 1], (m.start(), i)


def clean(text, items):
    m = re.search(r"\\begin\{document\}(.*)\\end\{document\}", text, re.S)
    if m:
        text = m.group(1)
    text = re.sub(r"^\\def.*$", "", text, flags=re.M)

    text = number_citations(text, items)
    _, block = extract_bibliography(text)
    if block:
        refs = "\n\n".join("[%d] %s" % (i, b) for i, (_, b) in enumerate(items, 1))
        text = text.replace(block, "\n\n" + refs + "\n\n")

    text, counters, labels = convert_tables_and_figures(text)

    # resolve \ref{...} to the printed table/figure number (the "Table"/"Fig." word
    # already precedes the reference in the source, so only the number is inserted)
    def ref_repl(m):
        kind, num = labels.get(m.group(1), ("Table", "?"))
        return "%s" % num

    text = re.sub(r"\\ref\{([^}]*)\}", ref_repl, text)

    # title and authors (placeholders must survive for the author to fill in)
    title, span = extract_command_arg(text, "title")
    if title is not None:
        title = re.sub(r"\\[a-zA-Z]+\s*", "", title).replace("{", "").replace("}", "").strip()
        text = text[:span[0]] + "\\begin{center}\\textbf{\\Large %s}\\end{center}\n" % title + text[span[1]:]
    author, span = extract_command_arg(text, "author")
    if author is not None:
        a = author
        a = re.sub(r"\\IEEEauthorblockN\{(.*?)\}", r"\1", a, flags=re.S)
        a = re.sub(r"\\IEEEauthorblockA\{(.*?)\}", r"\1", a, flags=re.S)
        a = a.replace("\\\\", ", ").replace("{", "").replace("}", "")
        a = re.sub(r"\\[a-zA-Z]+\s*", "", a)
        a = re.sub(r"\s+", " ", a).strip(" ,")
        text = text[:span[0]] + "\\begin{center}%s\\end{center}\n" % a + text[span[1]:]

    text = re.sub(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}",
                  lambda m: "\\n\\n\\textbf{Keywords:} " + m.group(1).strip() + "\\n\\n", text, flags=re.S)
    text = re.sub(r"\\begin\{abstract\}(.*?)\\end\{abstract\}",
                  lambda m: "\\n\\n\\textbf{Abstract.} " + m.group(1).strip() + "\\n\\n", text, flags=re.S)

    # captions become explicit numbered paragraphs
    text = re.sub(r"TABLECAP(\d+) ", lambda m: "\\textbf{Table %s.} " % m.group(1), text)
    text = re.sub(r"FIGCAP(\d+) ", lambda m: "\\textbf{Fig. %s.} " % m.group(1), text)

    # strip remaining LaTeX-only noise
    for pat in [r"\\label\{[^}]*\}", r"\\vspace\{[^}]*\}", r"\\centering", r"\\footnotesize",
                r"\\small", r"\\noindent", r"\\quad", r"\\hline", r"\\toprule",
                r"\\midrule", r"\\bottomrule", r"\\cmidrule\{[^}]*\}", r"\\setcounter\{[^}]*\}\{[^}]*\}",
                r"\\IEEEoverridecommandlockouts", r"\\maketitle"]:
        text = re.sub(pat, "", text)

    text = inline_math_to_text(text)
    return text, counters


def run_pandoc(src, out):
    exe = shutil.which("pandoc")
    if not exe:
        raise SystemExit("pandoc not found on PATH")
    cmd = [exe, "-f", "latex", "-t", "docx",
           "--resource-path", BUILD + os.pathsep + ROOT,
           "--wrap=none", src, "-o", out]
    p = subprocess.run(cmd, cwd=BUILD, capture_output=True, text=True, errors="replace")
    if p.returncode != 0 or not os.path.exists(out):
        raise SystemExit("pandoc failed:\n" + p.stdout[-1500:] + "\n" + p.stderr[-3000:])
    return out


def postprocess(docx_path, twocol=False):
    import docx
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document(docx_path)
    for section in d.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        sectPr = section._sectPr
        cols = sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectPr.append(cols)
        cols.set(qn("w:num"), "2" if twocol else "1")
        cols.set(qn("w:space"), "288")

    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rpr.set(qn(attr), "Times New Roman")
    pf = normal.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Inches(0.18)

    for p in d.paragraphs:
        txt = p.text.strip()
        is_caption = bool(re.match(r"^\*\*(Table|Fig\.)\s*\d+\.\*\*", txt)) or \
            bool(re.match(r"^(Table|Fig\.)\s*\d+\.", txt))
        style = p.style.name or ""
        is_heading = style.startswith("Heading")
        if is_heading:
            for r in p.runs:
                r.font.name = "Times New Roman"
                rpr = r._element.get_or_add_rPr().get_or_add_rFonts()
                for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                    rpr.set(qn(attr), "Times New Roman")
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.italic = True
                r.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            continue
        for r in p.runs:
            r.font.name = "Times New Roman"
            rpr = r._element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rpr.set(qn(attr), "Times New Roman")
            r.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
            if is_caption:
                r.font.size = Pt(8)
                r.italic = True
                p.alignment = 1
                p.paragraph_format.first_line_indent = Inches(0)
            elif r.font.size is None or r.font.size > Pt(10):
                r.font.size = Pt(10)
        if is_caption:
            p.paragraph_format.space_before = Pt(4)

    # tables: 8pt (matching the LaTeX \footnotesize), tight cell spacing, black text
    for t in d.tables:
        t.autofit = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        rpr = r._element.get_or_add_rPr().get_or_add_rFonts()
                        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                            rpr.set(qn(attr), "Times New Roman")
                        r.font.size = Pt(8)
                        r.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
    d.save(docx_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--twocol", action="store_true")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    stage_figures()
    text = open(TEX, encoding="utf-8").read()
    items, _ = extract_bibliography(text)
    cleaned, counters = clean(text, items)
    tmp = os.path.join(BUILD, "manuscript_docx.tex")
    open(tmp, "w", encoding="utf-8").write(cleaned)

    name = args.out or ("manuscript_2col.docx" if args.twocol else "manuscript.docx")
    out = os.path.join(ROOT, name)
    run_pandoc(tmp, out)
    postprocess(out, twocol=args.twocol)

    import docx
    d = docx.Document(out)
    body_text = " ".join([p.text for p in d.paragraphs] +
                          [c.text for t in d.tables for r in t.rows for c in r.cells])
    leftover = [p.text for p in d.paragraphs
                if re.search(r"\\(cite|ref|begin|end|textbf|emph|item|includegraphics)\b", p.text)]
    print("wrote %s" % out)
    print("  tables: %d (expected %d) | images: %d | references: %d | paragraphs: %d"
          % (len(d.tables), counters["table"], len(d.inline_shapes), len(items), len(d.paragraphs)))
    for cap in ("Table 1.", "Table 2.", "Table 3.", "Table 4.", "Fig. 1."):
        print("  caption %-9s %s" % (cap, "present" if cap in body_text else "MISSING"))
    print("  math converted to text:", "yes" if "$" not in body_text else "NO - raw $ left")
    if leftover:
        print("  WARNING: %d paragraphs still contain LaTeX commands, e.g. %r"
              % (len(leftover), leftover[0][:120]))


if __name__ == "__main__":
    sys.exit(main())
